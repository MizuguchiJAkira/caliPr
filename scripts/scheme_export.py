"""The measurements export for a study that follows another landmark scheme.

caliPr's traits are all defined on caliPr's own landmarks, so a study on a
published scheme (BGNN 2D, or one made here) has none of them, and mapping one
scheme's points onto the other's would be a claim about anatomy rather than a
conversion. What such a study does have, per specimen, is its landmarks. So its
measurements are those: one row per specimen, each landmark's x and y, in
millimetres when the specimen carries a scale and pixels when it does not, and
centroid size -- the size measure geometric morphometrics uses, defined on any
landmark set.

It goes through the same preview as the trait export -- one column per landmark,
with its count of specimens that lack it -- and honours the same choices: leave
landmarks out, keep only complete specimens, CSV or workbook.

Coordinates keep the photograph's orientation: origin at its top-left corner, x
to the right, y DOWN, as ImageJ measures. The .tps from the R export flips y, so
the two are mirror images; use one or the other. Every specimen here is in its
own photograph's frame, so the coordinates mean nothing across specimens until
they are superimposed (gpagen) -- which is what they are for.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import math
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT / "src"))
sys.path.insert(0, str(_ROOT / "scripts"))

from fish_morpho import grouping, schemes  # noqa: E402

from export_tps import px_per_mm  # noqa: E402


def scheme_of(dataset_dir: Path) -> dict | None:
    """The study's scheme, if it follows one other than caliPr's own."""
    f = dataset_dir / "schema.json"
    if not f.is_file():
        return None
    try:
        name = json.loads(f.read_text()).get("scheme")
    except Exception:
        return None
    return schemes.get(name)


def table(dataset_dir: Path, labels_dir: Path) -> tuple[list[str], dict, list[dict]]:
    """Every labelled specimen's landmarks, in the scheme's order.

    Returns (landmark order, landmark labels, rows); each row has fish_id, group,
    units, px_per_mm, data_note and ``points`` {name: (x, y) or None}.
    """
    order, labels = schemes.study_landmarks(dataset_dir)
    g_table = grouping.load_group_table(dataset_dir)
    g_pattern = grouping.filename_pattern(dataset_dir)
    rows = []
    for p in sorted(labels_dir.glob("*.json")):
        try:
            doc = json.loads(p.read_text())
        except Exception:
            continue
        fid = doc.get("fish_id") or p.stem
        kps = (doc.get("lateral") or {}).get("keypoints") or {}
        if not any(kps.get(n) for n in order):
            continue                                   # nothing of this scheme placed
        meta = doc.get("metadata") or {}
        ppm = px_per_mm(doc)
        pts = {n: ((kps[n][0] / ppm, kps[n][1] / ppm) if ppm else tuple(kps[n]))
               if kps.get(n) else None for n in order}
        rows.append({"fish_id": fid,
                     "group": grouping.resolve(fid, meta, g_table, g_pattern),
                     "units": "mm" if ppm else "px",
                     "px_per_mm": round(ppm, 4) if ppm else None,
                     "data_note": (meta.get("data_note") or "").replace("\n", " "),
                     "points": pts})
    rows.sort(key=lambda r: (r["group"] == "", r["group"], r["fish_id"]))
    return list(order), labels, rows


def centroid_size(pts: list[tuple[float, float]]) -> float:
    cx = sum(p[0] for p in pts) / len(pts)
    cy = sum(p[1] for p in pts) / len(pts)
    return math.sqrt(sum((p[0] - cx) ** 2 + (p[1] - cy) ** 2 for p in pts))


def export(dataset_dir: Path, labels_dir: Path, out_xlsx: Path, csv_path: Path | None = None,
           preview_json: Path | None = None, left_out: tuple[str, ...] = (),
           complete_only: bool = False) -> dict:
    """Write the workbook (and CSV and preview, when asked); return a summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    scheme = scheme_of(dataset_dir) or {}
    order, labels, rows = table(dataset_dir, labels_dir)
    kept = [n for n in order if n not in set(left_out)]
    incomplete = []
    if complete_only:
        incomplete = [r["fish_id"] for r in rows if any(r["points"][n] is None for n in kept)]
        rows = [r for r in rows if r["fish_id"] not in set(incomplete)]
        if not rows:
            raise RuntimeError("no specimen has every kept landmark — leave out more "
                               "landmarks, or export without 'complete rows only'")

    def csize(r):
        pts = [r["points"][n] for n in kept]
        return round(centroid_size(pts), 3) if kept and all(pts) else None

    meta_cols = ["fish_id", "group", "units", "px_per_mm", "centroid_size"]
    head = meta_cols + [f"{n}_{a}" for n in kept for a in ("x", "y")]

    def line(r):
        vals = [r["fish_id"], r["group"], r["units"], r["px_per_mm"], csize(r)]
        for n in kept:
            p = r["points"][n]
            vals += [round(p[0], 3), round(p[1], 3)] if p else [None, None]
        return vals

    wb = Workbook()
    ws = wb.active
    ws.title = "Coordinates"
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6E6E6")
    for r in rows:
        ws.append(line(r))
    for i, h in enumerate(head, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(h) + 2)

    names = wb.create_sheet("Landmarks")
    names.append(["index", "name", "label", "specimens without it"])
    for i, n in enumerate(order, start=1):
        names.append([i, n, labels.get(n, n),
                      "left out" if n in left_out else sum(r["points"][n] is None for r in rows)])

    mm = sum(r["units"] == "mm" for r in rows)
    about = wb.create_sheet("About")
    info = [
        ("caliPr landmark coordinates export", ""),
        ("", ""),
        ("dataset", dataset_dir.name),
        ("landmark scheme", f"{scheme.get('title', '?')}" +
                            (f" — {scheme['source']}" if scheme.get("source") else "")),
        ("generated", dt.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")),
        ("specimens", str(len(rows))),
        ("", ""),
        ("WHY COORDINATES", "caliPr's traits are defined on its own landmarks, so this "
                            "scheme has none. Its measurements are its landmarks: shape "
                            "is analysed from them after Procrustes superimposition."),
        ("Coordinates", "One row per specimen: each landmark's x and y. Origin at the "
                        "photograph's top-left corner, x to the right, y DOWN (as ImageJ "
                        "measures). The .tps from the R export flips y: use one or the other."),
        ("centroid_size", "Square root of the summed squared distances of the kept "
                          "landmarks from their centroid — the size measure of geometric "
                          "morphometrics. Blank where a kept landmark is missing."),
        ("Landmarks", "The scheme's landmarks in its numbered order, and how many "
                      "specimens lack each."),
        ("", ""),
        ("UNITS", ""),
        ("millimetres", f"{mm} specimen(s)"),
        ("pixels (no scale reference)", f"{len(rows) - mm} specimen(s)"),
    ]
    if left_out or complete_only:
        info += [("", ""), ("LEFT OUT AT EXPORT", ""),
                 ("landmarks", ", ".join(left_out) if left_out else "none")]
        if complete_only:
            info.append(("specimens", f"{len(incomplete)} missing a kept landmark: "
                                      + ", ".join(incomplete) if incomplete else "none"))
    for k, v in info:
        about.append([k, v])
    about["A1"].font = Font(bold=True, size=14)
    for r in range(1, about.max_row + 1):
        if about.cell(row=r, column=1).value in ("WHY COORDINATES", "UNITS", "LEFT OUT AT EXPORT"):
            about.cell(row=r, column=1).font = Font(bold=True)
        about.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    about.column_dimensions["A"].width = 30
    about.column_dimensions["B"].width = 100
    wb.move_sheet("About", offset=-(len(wb.sheetnames) - 1))
    wb.active = 0
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

    if csv_path:
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(head)
            for r in rows:
                w.writerow(["NA" if v is None and i >= 3 else ("" if v is None else v)
                            for i, v in enumerate(line(r))])

    if preview_json:
        # One preview column per landmark, the point shown as "x, y".
        preview_json.write_text(json.dumps({
            "what": "landmarks",
            "columns": [{"code": n, "label": labels.get(n, n),
                         "missing": sum(r["points"][n] is None for r in rows)} for n in kept],
            "rows": [{"id": r["fish_id"], "group": r["group"], "units": r["units"],
                      "values": [None if r["points"][n] is None else
                                 f"{r['points'][n][0]:.2f}, {r['points'][n][1]:.2f}"
                                 for n in kept]} for r in rows],
        }))
    return {"specimens": len(rows), "landmarks": len(kept), "mm": mm,
            "scheme": scheme.get("title", "?")}
