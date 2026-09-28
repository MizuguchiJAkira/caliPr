"""Measurement error: how much of a trait's variation is the person measuring it.

A difference between groups means little until the error of measuring it is
known. The standard way to find it is to measure the same specimens again,
blind, and split each trait's variance into the part between specimens and the
part between repeat measurements of one specimen (Bailey & Byrnes 1990;
Yezerinac, Lougheed & Handford 1992):

    %ME  = 100 * s2_within / (s2_within + s2_among)
    ICC(1) = s2_among / (s2_among + s2_within)            (repeatability;
                                                            Lessells & Boag 1987)

from a one-way ANOVA with specimen as the factor, where s2_within is the
within-specimen mean square and s2_among = (MS_among - MS_within) / k for k
measurements of each specimen. %ME is 100 * (1 - ICC(1)). Both count a
systematic shift between sessions -- one operator placing a landmark
consistently further forward than another -- as error. ICC(3,1), from the
two-way ANOVA with the sessions as raters (Shrout & Fleiss 1979), sets such a
constant shift aside: where it is well above ICC(1), the sessions differ
systematically, and the bias column says by how much.

The re-measurement is a *blind re-label round*: a study folder beside the
original (``<study>.relabel-<k>``) holding a stratified subset of the labelled
fish under codes (R1-01, R1-02 ...) in random order, with none of the original
landmarks and no model assistance. ``round.json`` there keeps which fish each
code is; the labeller never sees it. Rounds that re-label the same fish (a second
operator, a second week) are analysed together, as more sessions of one design.

Both sessions go through the same pipeline -- the same exclusions, calibration
and trait definitions -- so what differs between them is the labelling.
"""

from __future__ import annotations

import json
import math
from dataclasses import dataclass, field
from pathlib import Path

ROUND_FILE = "round.json"


# ---------------------------------------------------------------------------
# F distribution quantiles, for the confidence interval of ICC(1)
# ---------------------------------------------------------------------------

def _betacf(a: float, b: float, x: float) -> float:
    """Continued fraction for the incomplete beta function (modified Lentz)."""
    tiny, eps = 1e-300, 3e-15
    qab, qap, qam = a + b, a + 1.0, a - 1.0
    c, d = 1.0, 1.0 - qab * x / qap
    d = 1.0 / (d if abs(d) > tiny else tiny)
    h = d
    for m in range(1, 400):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1.0 + aa * d
        d = 1.0 / (d if abs(d) > tiny else tiny)
        c = 1.0 + aa / c
        c = c if abs(c) > tiny else tiny
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1.0 + aa * d
        d = 1.0 / (d if abs(d) > tiny else tiny)
        c = 1.0 + aa / c
        c = c if abs(c) > tiny else tiny
        delta = d * c
        h *= delta
        if abs(delta - 1.0) < eps:
            break
    return h


def _betainc(a: float, b: float, x: float) -> float:
    """Regularised incomplete beta I_x(a, b)."""
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    lbt = math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log1p(-x)
    if x < (a + 1) / (a + b + 2):
        return math.exp(lbt) * _betacf(a, b, x) / a
    return 1.0 - math.exp(lbt) * _betacf(b, a, 1 - x) / b


def f_cdf(f: float, d1: float, d2: float) -> float:
    if f <= 0:
        return 0.0
    return _betainc(d1 / 2, d2 / 2, d1 * f / (d1 * f + d2))


def f_ppf(p: float, d1: float, d2: float) -> float:
    """The F distribution's p-quantile, by bisection on its CDF."""
    lo, hi = 0.0, 1.0
    while f_cdf(hi, d1, d2) < p:
        hi *= 2
        if hi > 1e12:
            return math.inf
    for _ in range(200):
        mid = (lo + hi) / 2
        if f_cdf(mid, d1, d2) < p:
            lo = mid
        else:
            hi = mid
        if hi - lo < 1e-12 * max(1.0, hi):
            break
    return (lo + hi) / 2


# ---------------------------------------------------------------------------
# the statistics
# ---------------------------------------------------------------------------

@dataclass
class TraitError:
    """Measurement error of one trait over n specimens measured k times."""

    n: int
    k: int
    icc1: float | None = None
    icc1_lo: float | None = None
    icc1_hi: float | None = None
    icc31: float | None = None
    pct_me: float | None = None
    mean_abs_diff: float | None = None
    bias: float | None = None
    tem: float | None = None
    rel_tem: float | None = None
    mean: float | None = None

    @property
    def verdict(self) -> str:
        """Koo & Li (2016), on the lower confidence limit where there is one."""
        v = self.icc1_lo if self.icc1_lo is not None else self.icc1
        if v is None:
            return ""
        return ("excellent" if v >= 0.9 else "good" if v >= 0.75
                else "moderate" if v >= 0.5 else "poor")


def trait_error(rows: list[list[float]]) -> TraitError:
    """ICC, %ME and the rest for rows of k repeated measurements of a specimen.

    Only complete rows are used. Fewer than three specimens, or a trait that
    does not vary, leaves the coefficients empty rather than meaningless.
    """
    rows = [r for r in rows if r and all(v is not None and math.isfinite(v) for v in r)]
    n = len(rows)
    k = len(rows[0]) if rows else 0
    out = TraitError(n=n, k=k)
    if n < 3 or k < 2:
        return out
    grand = sum(sum(r) for r in rows) / (n * k)
    row_means = [sum(r) / k for r in rows]
    col_means = [sum(r[j] for r in rows) / n for j in range(k)]
    ss_rows = k * sum((m - grand) ** 2 for m in row_means)
    ss_cols = n * sum((m - grand) ** 2 for m in col_means)
    ss_total = sum((v - grand) ** 2 for r in rows for v in r)
    ss_within = ss_total - ss_rows
    ss_err = ss_total - ss_rows - ss_cols
    ms_rows = ss_rows / (n - 1)
    ms_within = ss_within / (n * (k - 1))
    ms_err = ss_err / ((n - 1) * (k - 1))
    out.mean = grand
    diffs = [r[j] - r[0] for r in rows for j in range(1, k)]
    out.mean_abs_diff = sum(abs(d) for d in diffs) / len(diffs)
    out.bias = sum(diffs) / len(diffs)
    # Technical error of measurement, generalised to k sessions: the root mean
    # within-specimen variance, in the trait's own units.
    out.tem = math.sqrt(ms_within)
    out.rel_tem = 100 * out.tem / grand if grand else None
    denom1 = ms_rows + (k - 1) * ms_within
    if denom1 <= 0:
        return out
    out.icc1 = (ms_rows - ms_within) / denom1
    s2_among = max((ms_rows - ms_within) / k, 0.0)
    out.pct_me = (100 * ms_within / (ms_within + s2_among)) if (ms_within + s2_among) > 0 else None
    denom31 = ms_rows + (k - 1) * ms_err
    out.icc31 = (ms_rows - ms_err) / denom31 if denom31 > 0 else None
    if ms_within > 0:
        F = ms_rows / ms_within
        d1, d2 = n - 1, n * (k - 1)
        fl = F / f_ppf(0.975, d1, d2)
        fu = F * f_ppf(0.975, d2, d1)
        out.icc1_lo = (fl - 1) / (fl + k - 1)
        out.icc1_hi = (fu - 1) / (fu + k - 1)
    else:
        out.icc1_lo = out.icc1_hi = 1.0
    return out


# ---------------------------------------------------------------------------
# rounds
# ---------------------------------------------------------------------------

def read_round(folder: Path) -> dict | None:
    f = Path(folder) / ROUND_FILE
    if not f.is_file():
        return None
    try:
        doc = json.loads(f.read_text())
    except Exception:
        return None
    return doc if isinstance(doc, dict) and doc.get("study") and doc.get("fish") else None


def rounds_of(study_dir: Path) -> list[dict]:
    """Every blind re-label round of this study, oldest first, with its folder."""
    study_dir = Path(study_dir)
    out = []
    for d in sorted(study_dir.parent.iterdir()):
        if not d.is_dir() or d == study_dir:
            continue
        r = read_round(d)
        if r and r["study"] == study_dir.name:
            r = dict(r, folder=d, done=sorted(done_codes(d)))
            out.append(r)
    out.sort(key=lambda r: (r.get("round", 0), r.get("created", "")))
    return out


def done_codes(round_dir: Path) -> set[str]:
    """Codes whose re-label has been saved with lateral landmarks."""
    sc = Path(round_dir) / "sidecars"
    done = set()
    for f in sc.glob("*.json") if sc.is_dir() else []:
        try:
            if ((json.loads(f.read_text()).get("lateral") or {}).get("keypoints")):
                done.add(f.stem)
        except Exception:
            pass
    return done


def operators_of(sidecar: dict) -> list[str]:
    """Who placed a sidecar's landmarks, the model left out, most frequent first."""
    from .operators import people
    return people(sidecar)


def _reference_fish(study_dir: Path) -> set[str]:
    """Fish shown as the worked example in the reference panel: re-labelling one
    of those blind would mean re-labelling it with its answers on screen."""
    out = set()
    ui = Path(__file__).resolve().parents[2] / "scripts" / "labeling_ui" / "reference.json"
    for f in (Path(study_dir) / "reference" / "reference.json", ui):
        try:
            sp = json.loads(f.read_text()).get("specimen")
            if sp:
                out.add(sp)
        except Exception:
            pass
    return out


def create_round(study_dir: Path, n: int, operator: str = "", same_as: int | None = None,
                 seed: int | None = None) -> dict:
    """Set up a blind re-label round of ``n`` labelled fish; returns its round.json.

    Fish are drawn across the study's comparison groups in turn, so a small round
    still covers every strain or population; within a group, at random. They get
    codes in random order, so the list gives away neither name nor group, and
    their photographs are linked in (a hard link, or a copy where the disk will
    not link), never moved. ``same_as`` re-labels the fish of an earlier round
    again -- a second operator, or the same one weeks later.
    """
    import datetime
    import os
    import random
    import shutil
    from . import grouping

    study_dir = Path(study_dir)
    if read_round(study_dir):
        raise ValueError("this is a re-label round already; start one from the study itself")
    existing = rounds_of(study_dir)
    k = max([r.get("round", 0) for r in existing] + [0]) + 1
    rng = random.Random(seed)
    if same_as is not None:
        prev = next((r for r in existing if r.get("round") == same_as), None)
        if prev is None:
            raise ValueError(f"there is no round {same_as} to repeat")
        chosen = sorted(prev["fish"].values())
    else:
        labelled = []
        for f in sorted((study_dir / "sidecars").glob("*.json")):
            try:
                doc = json.loads(f.read_text())
            except Exception:
                continue
            if (doc.get("lateral") or {}).get("keypoints") and _find_image(study_dir / "lateral", f.stem):
                labelled.append((f.stem, doc.get("metadata") or {}))
        skip = _reference_fish(study_dir)
        labelled = [(fid, m) for fid, m in labelled if fid not in skip]
        if len(labelled) < 3:
            raise ValueError("a round needs at least three labelled fish to re-label")
        table, pattern = grouping.load_group_table(study_dir), grouping.filename_pattern(study_dir)
        groups: dict[str, list[str]] = {}
        for fid, meta in labelled:
            groups.setdefault(grouping.resolve(fid, meta, table, pattern), []).append(fid)
        for g in groups.values():
            rng.shuffle(g)
        chosen, depth = [], 0
        n = max(3, min(int(n), len(labelled)))
        while len(chosen) < n:
            for g in sorted(groups):
                if depth < len(groups[g]) and len(chosen) < n:
                    chosen.append(groups[g][depth])
            depth += 1
    order = chosen[:]
    rng.shuffle(order)
    width = max(2, len(str(len(order))))
    fish = {f"R{k}-{i + 1:0{width}d}": fid for i, fid in enumerate(order)}

    folder = study_dir.parent / f"{study_dir.name}.relabel-{k}"
    folder.mkdir()
    (folder / "sidecars").mkdir()

    def link(src: Path, dst: Path):
        dst.parent.mkdir(parents=True, exist_ok=True)
        try:
            os.link(src, dst)
        except OSError:
            shutil.copy2(src, dst)
    for code, fid in fish.items():
        lat = _find_image(study_dir / "lateral", fid)
        link(lat, folder / "lateral" / f"{code}_L{lat.suffix}")
        fro_dir = study_dir / "frontal"
        if fro_dir.is_dir():
            for p in fro_dir.iterdir():
                if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".tif", ".tiff"} \
                        and p.stem in (fid, f"{fid}_F"):
                    link(p, folder / "frontal" / f"{code}_F{p.suffix}")
                    break
    for name in ("schema.json", "plausibility.json"):
        if (study_dir / name).is_file():
            shutil.copy2(study_dir / name, folder / name)
    doc = {
        "study": study_dir.name, "round": k, "blind": True,
        "created": datetime.datetime.now().astimezone().isoformat(timespec="seconds"),
        "operator": operator.strip(), "n": len(fish), "same_as": same_as, "seed": seed,
        "note": ("Blind re-label for measurement error. Which fish each code is stays in this "
                 "file; the labeller never sees it, nor the original landmarks. Label every "
                 "fish here from scratch, ruler included."),
        "fish": fish,
    }
    (folder / ROUND_FILE).write_text(json.dumps(doc, indent=2) + "\n")
    return dict(doc, folder=str(folder))


@dataclass
class Design:
    """Rounds that re-labelled the same fish, analysed together with the original."""

    rounds: list[dict]
    fish: list[str]
    sessions: list[str] = field(default_factory=list)       # descriptions, original first
    traits: dict[str, TraitError] = field(default_factory=dict)


def _find_image(folder: Path, fid: str) -> Path | None:
    """A fish's side-view photograph: named for it, or with the rig's _L suffix."""
    from .pipeline import IMAGE_EXTS
    if not Path(folder).is_dir():
        return None
    for p in sorted(Path(folder).iterdir()):
        if p.suffix.lower() in IMAGE_EXTS and p.stem in (fid, f"{fid}_L"):
            return p
    return None


def _values(spec, study_dir, excluded, group_table, group_pattern) -> dict[str, float]:
    from .pipeline import process_specimen
    rec = process_specimen(spec, group_table, group_pattern, excluded)
    return {code: mv.value for code, mv in rec.measurements.values.items()}


def analyse(study_dir: Path, labels_dir: Path | None = None, *, min_done: int = 3) -> list[Design]:
    """The measurement error of every trait, for each set of rounds on the same fish.

    A round counts once at least ``min_done`` of its fish are re-labelled; only
    fish re-labelled in every round of a design are used.
    """
    from . import grouping
    from .pipeline import SpecimenInput, _excluded_structures
    study_dir = Path(study_dir)
    labels_dir = Path(labels_dir) if labels_dir else study_dir / "sidecars"
    excluded = _excluded_structures(study_dir)
    g_table = grouping.load_group_table(study_dir)
    g_pattern = grouping.filename_pattern(study_dir)

    by_fish: dict[frozenset, list[dict]] = {}
    for r in rounds_of(study_dir):
        if len(r["done"]) >= min_done:
            by_fish.setdefault(frozenset(r["fish"].values()), []).append(r)

    designs = []
    for rounds in by_fish.values():
        # fish every round of this design has re-labelled, and the original has
        done = set(rounds[0]["fish"].values())
        for r in rounds:
            done &= {r["fish"][c] for c in r["done"] if c in r["fish"]}
        fish = sorted(f for f in done if (labels_dir / f"{f}.json").is_file())
        if len(fish) < min_done:
            continue
        values: dict[str, list[dict[str, float]]] = {}
        orig_ops: set[str] = set()
        round_ops: list[set[str]] = [set() for _ in rounds]
        for f in fish:
            sc = json.loads((labels_dir / f"{f}.json").read_text())
            orig_ops.update(operators_of(sc))
            img = _find_image(study_dir / "lateral", f)
            row = [_values(SpecimenInput(f, img, labels_dir / f"{f}.json", sc),
                           study_dir, excluded, g_table, g_pattern)]
            for i, r in enumerate(rounds):
                code = next(c for c, x in r["fish"].items() if x == f)
                rp = r["folder"] / "sidecars" / f"{code}.json"
                rsc = json.loads(rp.read_text())
                round_ops[i].update(operators_of(rsc))
                rimg = _find_image(r["folder"] / "lateral", code) or img
                row.append(_values(SpecimenInput(code, rimg, rp, rsc),
                                   r["folder"], excluded, g_table, g_pattern))
            values[f] = row
        d = Design(rounds=rounds, fish=fish)
        d.sessions = ["original labels" + (f" (operator {', '.join(sorted(orig_ops))})"
                                           if orig_ops else " (operator not recorded)")]
        for r, ops in zip(rounds, round_ops):
            who = ", ".join(sorted(ops)) or r.get("operator") or "not recorded"
            d.sessions.append(f"blind re-label round {r.get('round')} ({r.get('created', '')[:10]}, "
                              f"operator {who})")
        codes = sorted({c for row in values.values() for s in row for c in s})
        for c in codes:
            rows = [[s.get(c) for s in values[f]] for f in fish]
            rows = [[v if (v is not None and not (isinstance(v, float) and math.isnan(v))) else None
                     for v in r] for r in rows]
            d.traits[c] = trait_error(rows)
        designs.append(d)
    return designs


# ---------------------------------------------------------------------------
# the workbook sheet
# ---------------------------------------------------------------------------

REFERENCES = (
    "Bailey RC, Byrnes J (1990) A new, old method for assessing measurement error in both "
    "univariate and multivariate morphometric studies. Syst Zool 39:124-130.",
    "Yezerinac SM, Lougheed SC, Handford P (1992) Measurement error and morphometric studies: "
    "statistical power and observer experience. Syst Biol 41:471-482.",
    "Lessells CM, Boag PT (1987) Unrepeatable repeatabilities: a common mistake. Auk 104:116-121.",
    "Shrout PE, Fleiss JL (1979) Intraclass correlations: uses in assessing rater reliability. "
    "Psychol Bull 86:420-428.",
    "Koo TK, Li MY (2016) A guideline of selecting and reporting intraclass correlation "
    "coefficients for reliability research. J Chiropr Med 15:155-163.",
)


def add_sheet(workbook: Path, study_dir: Path, labels_dir: Path | None = None,
              keep: list[str] | None = None) -> dict:
    """Add the Measurement error sheet, when the study has a round to report.

    ``keep`` is the workbook's own trait columns, so the sheet reports exactly
    the traits the file does. Returns what was added, for the caller to say.
    """
    designs = analyse(study_dir, labels_dir)
    if not designs:
        return {"designs": 0}
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from .landmark_config import TRAITS
    by_code = {t.code: t for t in TRAITS}
    wb = load_workbook(workbook)
    if keep is None:
        ws0 = wb["Measurements"]
        hdr = [str(c.value) for c in ws0[1]]
        keep = [h.split(" — ")[0] for h in hdr[hdr.index("units") + 1:]] if "units" in hdr else None
    ws = wb.create_sheet("Measurement error")
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E6E6E6")
    ws.append(["Measurement error", ""])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["", "The same fish labelled again, blind -- none of the original landmarks shown, "
                   "no model assistance -- and every trait measured both times by the same "
                   "pipeline. %ME is the share of a trait's variance that is measurement error; "
                   "ICC(1) is repeatability (= 1 - %ME/100). ICC(3,1) sets aside a constant shift "
                   "between sessions: well above ICC(1) means the sessions differ systematically "
                   "(see bias). Verdicts follow Koo & Li (2016), judged on the lower 95% limit "
                   "of ICC(1)."])
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 64
    head = ["trait", "name", "unit", "fish", "sessions", "ICC(1)", "95% CI low", "95% CI high",
            "%ME", "ICC(3,1)", "mean |re-label - original|", "bias (re-label - original)",
            "TEM", "relative TEM %", "verdict"]
    added = 0
    for d in designs:
        ws.append([])
        ws.append([f"{len(d.fish)} fish"] + [""] * 0)
        ws.cell(row=ws.max_row, column=1).font = bold
        for i, s in enumerate(d.sessions):
            ws.append([f"session {i + 1}", s])
        ws.append(head)
        for c in ws[ws.max_row]:
            c.font = bold
            c.fill = fill
        for code, te in d.traits.items():
            if keep is not None and code not in keep:
                continue
            t = by_code.get(code)
            unit = {"mm": "mm", "mm^2": "mm²", "deg": "°"}.get(t.unit.value, "") if t else ""
            r3 = (lambda v: None if v is None else round(v, 3))
            ws.append([code, t.label if t else code, unit, te.n, te.k, r3(te.icc1), r3(te.icc1_lo),
                       r3(te.icc1_hi), None if te.pct_me is None else round(te.pct_me, 1),
                       r3(te.icc31), r3(te.mean_abs_diff), r3(te.bias), r3(te.tem),
                       None if te.rel_tem is None else round(te.rel_tem, 2), te.verdict])
            added += 1
    ws.append([])
    ws.append(["METHOD", "One-way ANOVA with specimen as factor: %ME = 100 s²within / (s²within + "
                         "s²among), s²among = (MSamong - MSwithin)/k (Bailey & Byrnes 1990; "
                         "Yezerinac et al. 1992). ICC(1) = (MSamong - MSwithin)/(MSamong + "
                         "(k-1)MSwithin) with an F-based 95% interval; ICC(3,1) = (MSspecimen - "
                         "MSerror)/(MSspecimen + (k-1)MSerror) from the two-way ANOVA with sessions "
                         "as raters (Shrout & Fleiss 1979). TEM = √MSwithin, "
                         "in the trait's units. Values in pixels where the photograph has no scale."])
    ws.cell(row=ws.max_row, column=1).font = bold
    ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[ws.max_row].height = 64
    for ref in REFERENCES:
        ws.append(["", ref])
    for col, w in zip("ABCDEFGHIJKLMNO", (12, 60, 6, 6, 9, 9, 11, 11, 8, 9, 14, 14, 9, 11, 11)):
        ws.column_dimensions[col].width = w
    if "About" in wb.sheetnames:
        a = wb["About"]
        a.append([])
        a.append(["MEASUREMENT ERROR", f"{sum(len(d.fish) for d in designs)} fish re-labelled "
                                       f"blind in {sum(len(d.rounds) for d in designs)} round(s): "
                                       "ICC and %ME per trait on the Measurement error sheet."])
        a.cell(row=a.max_row, column=1).font = bold
    wb.save(workbook)
    return {"designs": len(designs), "traits": added,
            "fish": sum(len(d.fish) for d in designs)}
