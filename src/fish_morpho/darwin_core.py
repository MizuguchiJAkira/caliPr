"""Darwin Core records for a study's specimens.

A landmarked photograph is only as useful as the specimen it can be traced back
to. A museum lot is identified by the institution that holds it and its catalogue
number; the fish in it by its scientific name, where and when it was collected,
and by whom. Darwin Core (https://dwc.tdwg.org/terms/) is the vocabulary every
collection database, iDigBio and GBIF use for exactly those facts, so recording
them under its term names means a measurement can be joined back to the museum's
own record, merged with another lab's, and credited.

Each study keeps them in one table, ``darwin_core.csv`` in the study folder: a
``fish_id`` column, then one column per term below. A table rather than the
sidecars, because the records exist before any fish is labelled -- they come
from the collection, often for a whole lot at once -- and because a spreadsheet
is how they arrive.

Three ways in, all explicit:

* **Typed** in the labeler's Specimen records table.
* **Imported** from a CSV or Excel sheet: an iDigBio harvest, a museum's export,
  a lab's data sheet. Its columns are matched to terms by name, with the usual
  spellings (``catalog_number``, ``Collector``, genus and species as separate
  columns) understood; its rows are matched to photographs by id or filename --
  or, for a museum's table with one row per lot, by catalogue number, reaching
  every photograph of that lot.
* **Suggested** from the filenames -- ``1947_CUMV_68133_02`` holds an institution
  and a catalogue number, ``Salvelinus_fontinalis_ASN_31`` a scientific name --
  and written only when accepted. As with comparison groups (see
  :mod:`fish_morpho.grouping`), a value nobody chose never reaches an export.

The exports carry them: the measurements workbook gains a Specimens sheet and a
MeasurementOrFact sheet (Darwin Core's own shape for traits: one row per value,
with its unit and the definition it was measured by), and the files for R carry
the identifiers beside each specimen.

Landmark coordinates are not a Darwin Core term and are not forced into one;
they stay in the TPS and per-specimen files, joined to these records by fish_id.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

FILENAME = "darwin_core.csv"
KEY = "fish_id"

#: The terms kept, in column order, each with what goes in it. A small subset of
#: Darwin Core: what identifies a specimen, what it is, and where it came from.
TERMS: tuple[tuple[str, str], ...] = (
    ("institutionCode", "The collection's institution, as the collection writes it: CUMV, MCZ, USNM."),
    ("collectionCode", "The collection within it, e.g. Fish."),
    ("catalogNumber", "The lot's catalogue number. Every fish photographed from one lot shares it."),
    ("occurrenceID", "The globally unique id the collection publishes for the record, if it has one "
                     "(a URN or a URL). Copied, never made up."),
    ("scientificName", "Genus and species, e.g. Salvelinus fontinalis."),
    ("basisOfRecord", "PreservedSpecimen, LivingSpecimen, MaterialSample or HumanObservation."),
    ("sex", "female, male or undetermined."),
    ("lifeStage", "adult, juvenile, larva ..."),
    ("fieldNumber", "The number written in the field, e.g. ST-HRN LMH 2/24/2025."),
    ("recordedBy", "Who collected it."),
    ("eventDate", "When it was collected, ISO 8601: 1947, 1947-06 or 1947-06-12."),
    ("country", "Country."),
    ("stateProvince", "State or province."),
    ("county", "County."),
    ("locality", "Where, in words."),
    ("associatedMedia", "Where the source photograph is published (URL)."),
    ("license", "The licence the record or photograph is shared under."),
)
TERM_NAMES: tuple[str, ...] = tuple(t for t, _ in TERMS)

# How other tables spell the same thing. Keys are lower-case with everything but
# letters and digits removed, so "Catalog No." and "catalog_number" both match.
_ALIASES: dict[str, str] = {
    **{t.lower(): t for t in TERM_NAMES},
    "institution": "institutionCode", "inst": "institutionCode", "museum": "institutionCode",
    "collection": "collectionCode",
    "catalog": "catalogNumber", "catalogno": "catalogNumber", "catno": "catalogNumber",
    "catalognumber": "catalogNumber", "cataloguenumber": "catalogNumber",
    "catalogueno": "catalogNumber", "lot": "catalogNumber", "lotnumber": "catalogNumber",
    "scientificname": "scientificName", "taxon": "scientificName", "binomial": "scientificName",
    "species": "scientificName",
    "sex": "sex", "gender": "sex",
    "lifestage": "lifeStage", "stage": "lifeStage", "maturity": "lifeStage",
    "fieldno": "fieldNumber", "fieldnumber": "fieldNumber",
    "collector": "recordedBy", "collectors": "recordedBy", "collectedby": "recordedBy",
    "recordedby": "recordedBy",
    "date": "eventDate", "year": "eventDate", "collectiondate": "eventDate",
    "datecollected": "eventDate", "eventdate": "eventDate",
    "state": "stateProvince", "province": "stateProvince", "stateprovince": "stateProvince",
    "site": "locality",
    "imageurl": "associatedMedia", "mediaurl": "associatedMedia",
    "licence": "license",
}

#: Columns that can say which photograph a row is about, most specific first.
_KEY_COLUMNS = ("fishid", "id", "specimen", "specimenid", "localfilename", "filename",
                "file", "imagefilename", "image", "photo", "label")


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(name or "").lower())


def path(study_dir: Path) -> Path:
    return Path(study_dir) / FILENAME


def load(study_dir: Path) -> dict[str, dict[str, str]]:
    """{fish_id: {term: value}} with blank values left out; empty if no table.

    A malformed table is not worth failing an export over, as with groups.csv:
    the records are additional information, and losing them degrades the export
    rather than invalidating it. Columns that are not kept terms are ignored.
    """
    f = path(study_dir)
    if not f.is_file():
        return {}
    try:
        rows = list(csv.DictReader(io.StringIO(f.read_text(encoding="utf-8-sig"))))
    except Exception:
        return {}
    out: dict[str, dict[str, str]] = {}
    for r in rows:
        fid = str(r.get(KEY) or "").strip()
        if not fid:
            continue
        rec = {t: str(r.get(t) or "").strip() for t in TERM_NAMES}
        rec = {t: v for t, v in rec.items() if v}
        if rec:
            out[fid] = rec
    return out


def save(study_dir: Path, records: dict[str, dict[str, str]]) -> Path:
    """Write the table: sorted by fish_id, every kept term a column, blanks blank.

    A fish with nothing recorded has no row. Written whole, through a temporary
    file, so a crash mid-write cannot leave half a table.
    """
    f = path(study_dir)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([KEY, *TERM_NAMES])
    for fid in sorted(records):
        rec = records[fid] or {}
        if any(str(rec.get(t) or "").strip() for t in TERM_NAMES):
            w.writerow([fid, *(str(rec.get(t) or "").strip() for t in TERM_NAMES)])
    tmp = f.with_suffix(".csv.tmp")
    tmp.write_text(buf.getvalue(), encoding="utf-8")
    tmp.replace(f)
    return f


def update(study_dir: Path, changes: dict[str, dict[str, str]]) -> dict[str, dict[str, str]]:
    """Apply {fish_id: {term: value}}: a value sets the term, "" clears it.

    Terms not named are left as they were; names that are not kept terms are
    refused rather than silently dropped.
    """
    bad = sorted({t for c in changes.values() for t in (c or {}) if t not in TERM_NAMES})
    if bad:
        raise ValueError(f"not a Darwin Core term kept here: {', '.join(bad)}")
    records = load(study_dir)
    for fid, change in changes.items():
        rec = dict(records.get(fid) or {})
        for t, v in (change or {}).items():
            v = str(v if v is not None else "").strip()
            if v:
                rec[t] = v
            else:
                rec.pop(t, None)
        if rec:
            records[fid] = rec
        else:
            records.pop(fid, None)
    save(study_dir, records)
    return records


# ---------------------------------------------------------------------------
# Suggestions from filenames
# ---------------------------------------------------------------------------

_CUMV = re.compile(r"(?:^|_)(CUMV)[A-Za-z]*_(\d+)", re.IGNORECASE)
_BINOMIAL = re.compile(r"^([A-Z][a-z]+)_([a-z]{3,})(?:_|$)")


def suggest(fish_id: str, study_name: str = "") -> dict[str, str]:
    """What the filename (or, failing that, the study's name) says, if anything.

    Only what is written there: CUMV and a lot number, a leading Genus_species.
    A leading year is not read as a collection date, because nothing in the
    filename says which date it is.
    """
    out: dict[str, str] = {}
    m = _CUMV.search(fish_id) or _CUMV.search(study_name or "")
    if m:
        out["institutionCode"] = m.group(1).upper()
        out["catalogNumber"] = m.group(2)
    b = _BINOMIAL.match(fish_id)
    if b:
        out["scientificName"] = f"{b.group(1)} {b.group(2)}"
    return out


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def read_table(name: str, data: bytes) -> list[dict[str, str]]:
    """Rows of a CSV or Excel upload as {header: value}, header row first."""
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip() for h in next(it, [])]
        rows = []
        for r in it:
            if r is None or all(v in (None, "") for v in r):
                continue
            rows.append({h: _cell(v) for h, v in zip(head, r) if h})
        return rows
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    return [{str(k or "").strip(): _cell(v) for k, v in r.items() if k}
            for r in csv.DictReader(io.StringIO(text), dialect=dialect)]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))                 # Excel reads a catalogue number as 68133.0
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return str(v).strip()


def _stems(value: str) -> set[str]:
    """Ways a key cell can name a photograph: as written, without its extension,
    and without a trailing _L / _F view suffix."""
    v = str(value or "").strip()
    stem = re.sub(r"\.[A-Za-z0-9]{2,5}$", "", Path(v).name)
    out = {v, stem, re.sub(r"_[LF]$", "", stem)}
    return {s.lower() for s in out if s}


def match_columns(headers) -> dict[str, str]:
    """{column in the upload: term}. Genus + species columns are joined later."""
    out: dict[str, str] = {}
    have = {_norm(h) for h in headers}
    for h in headers:
        n = _norm(h)
        if n == "species" and "genus" in have:
            continue                        # joined with genus into scientificName
        t = _ALIASES.get(n)
        if t and t not in out.values():
            out[h] = t
    return out


def _catkey(value) -> str:
    """A catalogue number reduced to what two tables agree on: "CUMV 68133",
    "CUMV-68133" and "068133" are all 68133."""
    v = str(value or "").strip()
    m = re.search(r"(\d+)\s*$", v)
    return (m.group(1).lstrip("0") or "0") if m else _norm(v)


def plan_import(rows: list[dict[str, str]], fish_ids, existing=None) -> dict:
    """What importing these rows would change, without changing anything.

    Rows are matched to photographs one of two ways, whichever reaches more of
    them. By a column naming the photograph -- fish_id, a filename -- one row
    per photograph. Or, for a museum's own table, which has one row per lot, by
    catalogue number: a row reaches every photograph whose record already has
    that catalogueNumber (and the same institution, where both say). Other
    columns are mapped to terms by name. Returns::

        {"key": column, "by": "photo" | "lot", "columns": {column: term},
         "changes": {fish_id: {term: value}},
         "named": photographs the rows reach, "matched": of those, how many gain a value,
         "unmatched": [row keys reaching no photograph],
         "unused": [columns that are not terms], "overwrites": n}

    Blank cells change nothing: an import adds and corrects, it never erases.
    """
    existing = existing or {}
    fish_ids = list(fish_ids)
    lookup: dict[str, str] = {}
    for fid in fish_ids:
        for s in _stems(fid):
            lookup.setdefault(s, fid)
    headers = list(rows[0].keys()) if rows else []

    def by_photo(h):
        def reach(r):
            for s in _stems(r.get(h)):
                if s in lookup:
                    return [lookup[s]]
            return []
        return reach

    def reached(reach) -> int:
        return len({f for r in rows for f in reach(r)})

    best, reach, by, best_n = None, (lambda r: []), "photo", 0
    named_cols = [h for want in _KEY_COLUMNS for h in headers if _norm(h) == want]
    for pool in (named_cols, headers):       # a named key column first, then any
        for h in pool:
            n = reached(by_photo(h))
            if n > best_n:
                best, reach, best_n = h, by_photo(h), n
        if best is not None:
            break

    # A museum's table: one row per lot, reaching every photograph of it.
    cat = next((h for h in headers if _ALIASES.get(_norm(h)) == "catalogNumber"), None)
    inst = next((h for h in headers if _ALIASES.get(_norm(h)) == "institutionCode"), None)
    lots: dict[str, list[tuple[str, str]]] = {}
    for fid in fish_ids:
        rec = existing.get(fid) or {}
        if rec.get("catalogNumber"):
            lots.setdefault(_catkey(rec["catalogNumber"]), []).append(
                (fid, _norm(rec.get("institutionCode"))))
    if cat and lots and cat != best:
        def by_lot(r):
            k = _catkey(r.get(cat))
            i = _norm(r.get(inst)) if inst else ""
            return [f for f, fi in lots.get(k, []) if not (i and fi and i != fi)] if k else []
        n = reached(by_lot)
        if n > best_n:
            best, reach, by, best_n = cat, by_lot, "lot", n

    cols = match_columns([h for h in headers if h != best])
    if by == "lot" and inst in cols:
        del cols[inst]                       # it chose the lot; it is not news
    genus = next((h for h in headers if _norm(h) == "genus"), None)
    species = next((h for h in headers if _norm(h) == "species"), None)
    joined = bool(genus and species and "scientificName" not in cols.values())

    changes: dict[str, dict[str, str]] = {}
    unmatched: list[str] = []
    named: set[str] = set()
    for r in rows:
        fids = reach(r) if best else []
        if not fids:
            if best and str(r.get(best) or "").strip():
                unmatched.append(str(r.get(best)).strip())
            continue
        named.update(fids)
        rec: dict[str, str] = {}
        for h, t in cols.items():
            v = str(r.get(h) or "").strip()
            if t == "scientificName" and _norm(h) == "species" and " " not in v:
                continue                    # an epithet alone is not a name
            if v:
                rec[t] = v
        if joined:
            g, sp = str(r.get(genus) or "").strip(), str(r.get(species) or "").strip()
            if g and sp:
                rec["scientificName"] = f"{g} {sp}" if not sp.startswith(g) else sp
        for fid in fids:
            if rec:
                changes.setdefault(fid, {}).update(rec)
    overwrites = sum(1 for fid, rec in changes.items() for t, v in rec.items()
                     if (existing.get(fid) or {}).get(t) and existing[fid][t] != v)
    used = set(cols) | ({genus, species} if joined else set()) | {best}
    if by == "lot" and inst:
        used.add(inst)
    return {"key": best, "by": by,
            "columns": {**cols, **({f"{genus} + {species}": "scientificName"} if joined else {})},
            "changes": changes, "named": len(named), "matched": len(changes),
            "unmatched": unmatched,
            "unused": [h for h in headers if h not in used], "overwrites": overwrites}


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------

def used_terms(records: dict[str, dict[str, str]], fish_ids=None) -> list[str]:
    """The kept terms with a value for at least one of these fish, in term order."""
    ids = set(fish_ids) if fish_ids is not None else set(records)
    return [t for t in TERM_NAMES if any((records.get(f) or {}).get(t) for f in ids)]


_UNIT = {"mm": "mm", "mm^2": "mm2", "deg": "degrees"}


def add_sheets(workbook: Path, study_dir: Path, *, measurements: bool = True,
               labels_dir: Path | None = None) -> dict:
    """Add the Specimens sheet (and, for a trait workbook, MeasurementOrFact).

    Read back from the workbook's own first data sheet, so the records are the
    specimens and values the file actually holds. Returns counts for the caller
    to report. A study with no records gets nothing added: a workbook is not
    padded with a sheet of blanks.
    """
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    records = load(study_dir)
    if not records:
        return {"specimens": 0, "with_record": 0, "terms": [], "measurement_rows": 0}
    wb = load_workbook(workbook)
    first = wb["Measurements"] if "Measurements" in wb.sheetnames else wb[
        next(n for n in wb.sheetnames if n != "About")]
    it = first.iter_rows(values_only=True)
    hdr = [str(h) if h is not None else "" for h in next(it)]
    rows = [r for r in it if r and r[0] not in (None, "")]
    ids = [str(r[hdr.index(KEY)]) for r in rows]
    terms = used_terms(records, ids)
    with_record = sum(1 for f in ids if records.get(f))

    def style(ws):
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E6E6E6")
        ws.freeze_panes = "B2"

    ws = wb.create_sheet("Specimens")
    ws.append([KEY, *terms] if terms else [KEY, "(no Darwin Core records for this study yet)"])
    for f in ids:
        ws.append([f, *((records.get(f) or {}).get(t, "") for t in terms)])
    style(ws)
    ws.column_dimensions["A"].width = 34
    for i, t in enumerate(terms, start=2):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(t) + 4)

    mof = 0
    if measurements and "units" in hdr:
        from .landmark_config import TRAITS
        by_code = {t.code: t for t in TRAITS}
        u = hdr.index("units")
        from .operators import determined_by
        labels = Path(labels_dir) if labels_dir else Path(study_dir) / "sidecars"
        sidecars: dict[str, dict] = {}
        for f in ids:
            try:
                sidecars[f] = json.loads((labels / f"{f}.json").read_text())
            except Exception:
                sidecars[f] = {}
        ws = wb.create_sheet("MeasurementOrFact")
        ws.append([KEY, "measurementType", "measurementValue", "measurementUnit",
                   "measurementDeterminedBy", "measurementMethod", "measurementRemarks"])
        for r, f in zip(rows, ids):
            unit_row = r[u] or ""
            for h, v in zip(hdr[u + 1:], r[u + 1:]):
                if v in (None, ""):
                    continue
                code = h.split(" — ")[0]
                t = by_code.get(code)
                name = f"{t.label} ({code})" if t else code
                unit = (_UNIT.get(t.unit.value, t.unit.value) if t else "")
                if unit_row == "px" and unit in ("mm", "mm2"):
                    unit = "px" if unit == "mm" else "px2"
                method = (f"caliPr, from landmarks on a photograph: {t.description}"
                          if t else "caliPr")
                who = determined_by(sidecars.get(f) or {}, (*t.required_keypoints,
                                                           *t.required_polygons)) if t else ""
                ws.append([f, name, v, unit, who, method,
                           "no scale in the photograph; pixels" if unit_row == "px" else ""])
                mof += 1
        style(ws)
        for col, w in zip("ABCDEFG", (34, 34, 16, 16, 18, 60, 30)):
            ws.column_dimensions[col].width = w

    if "About" in wb.sheetnames:
        a = wb["About"]
        a.append([])
        a.append(["DARWIN CORE", ""])
        a["A" + str(a.max_row)].font = Font(bold=True)
        a.append(["Specimens", f"Darwin Core records (https://dwc.tdwg.org/terms/) for "
                               f"{with_record} of {len(ids)} specimens, from {FILENAME} "
                               f"in the study folder. Blank where nothing was recorded."])
        if measurements and "units" in hdr:
            a.append(["MeasurementOrFact", "The Measurements sheet in Darwin Core's own "
                                           "long form: one row per value, with its unit "
                                           "and the definition it was measured by."])
    wb.save(workbook)
    return {"specimens": len(ids), "with_record": with_record, "terms": terms,
            "measurement_rows": mof}
